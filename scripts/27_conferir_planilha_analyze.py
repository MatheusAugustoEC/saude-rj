"""Confere a planilha recalculada do Analyze contra o Python (numpy): erros de fórmula e diferenças de valor.

Uso: python -X utf8 scripts/27_conferir_planilha_analyze.py
"""
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

import _mapa_comum as c

RAIZ = Path(__file__).resolve().parent.parent
wb = load_workbook(RAIZ / "planilhas" / "saude_fiscal_rj_analyze_calc.xlsx", data_only=True)
ERROS = {"#DIV/0!", "#NAME?", "#VALUE!", "#REF!", "#N/A", "#NUM!", "#NULL!"}
erros = [(w.title, x.coordinate) for w in wb for row in w.iter_rows() for x in row if isinstance(x.value, str) and x.value in ERROS]
print("erros de fórmula:", len(erros), erros[:5])

mun, _, _ = c.carrega_dados()
mun = mun.assign(ord=(mun.grupo == c.CAPITAL).astype(int)).sort_values(["ord", "ente"]).reset_index(drop=True)
g, rc, fa = mun.saude / mun.populacao, mun.receita / mun.populacao, mun.saude / mun.receita
lg, lr, lf = np.log(g), np.log(rc), np.log(fa)
sem = mun.grupo != c.CAPITAL
cap = mun[mun.grupo == c.CAPITAL].index[0]
esperado = {
    "gasto por habitante (R$)": float(np.exp(lg.mean())), "receita por habitante (R$)": float(np.exp(lr.mean())),
    "fatia da receita na saúde": float(np.exp(lf.mean())),
    "gasto x receita": float(np.corrcoef(lg, lr)[0, 1]), "gasto x fatia": float(np.corrcoef(lg, lf)[0, 1]),
    "elasticidade b, 92 cidades": float(np.polyfit(lr, lg, 1)[0]), "R2, 92 cidades": float(np.corrcoef(lg, lr)[0, 1] ** 2),
    "elasticidade b, sem a capital (91)": float(np.polyfit(lr[sem], lg[sem], 1)[0]),
    "R2, sem a capital (91)": float(np.corrcoef(lg[sem], lr[sem])[0, 1] ** 2),
    "gasto por habitante da capital (R$)": float(g[cap]), "receita por habitante da capital (R$)": float(rc[cap]),
    "fatia da capital": float(fa[cap]),
    "gasto da capital vs referência": float(np.exp(lg[cap] - lg.mean()) - 1),
    "receita da capital vs referência": float(np.exp(lr[cap] - lr.mean()) - 1),
    "fatia da capital vs referência": float(np.exp(lf[cap] - lf.mean()) - 1),
    "mediana da fatia nas 92": float(fa.median()),
}
# modelo M3 (91 cidades, sem a capital) por mínimos quadrados e previsão da capital com intervalo de previsão de 95%
from scipy import stats

X = np.column_stack([np.ones(sem.sum()), lr[sem], (mun.grupo[sem] == c.BAIXADA).astype(float), (mun.grupo[sem] == c.RESTO).astype(float)])
yv = lg[sem].values
beta, rss, *_ = np.linalg.lstsq(X, yv, rcond=None)
n, k = X.shape
s = float(np.sqrt(((yv - X @ beta) ** 2).sum() / (n - k)))
x0 = np.array([1, lr[cap], 0, 1.0])
h = float(x0 @ np.linalg.inv(X.T @ X) @ x0)
tcrit = stats.t.ppf(0.975, n - k)
prev = float(x0 @ beta)
esperado.update({
    "M3 elasticidade b": float(beta[1]), "M3 efeito Baixada (%)": float(np.exp(beta[2]) - 1), "M3 efeito Resto da Metropolitana (%)": float(np.exp(beta[3]) - 1),
    "M3 R2": float(1 - ((yv - X @ beta) ** 2).sum() / ((yv - yv.mean()) ** 2).sum()), "M3 erro-padrão da regressão (s)": s,
    "gasto previsto (R$ por habitante)": float(np.exp(prev)), "alavanca da capital (x0 (X'X)^-1 x0')": h,
    "limite inferior do intervalo de previsão (R$)": float(np.exp(prev - tcrit * s * np.sqrt(1 + h))),
    "limite superior do intervalo de previsão (R$)": float(np.exp(prev + tcrit * s * np.sqrt(1 + h))),
    "gasto real da capital em relação ao previsto": float(g[cap] / np.exp(prev) - 1),
})
resumo = {row[0].value: row[1].value for row in wb["Resumo"].iter_rows(min_row=3) if row[1].value is not None}
dif = 0
for rot, e in esperado.items():
    val = resumo[rot]
    ok = abs(val - e) <= 1e-6 * max(1, abs(e))
    dif += not ok
    print(f"{'OK  ' if ok else 'DIF '} {rot}: planilha {val:.6f} | python {e:.6f}")
print("maior erro da identidade na planilha:", resumo["maior erro da identidade (deve ser ~0)"])
print("diferenças:", dif)

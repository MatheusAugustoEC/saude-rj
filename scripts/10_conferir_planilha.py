"""Confere a planilha recalculada contra o notebook: erros de fórmula e diferenças de valor.

Uso: python -u -I scripts/10_conferir_planilha.py [caminho_da_planilha]
Padrão: planilhas/saude_fiscal_rj_2024_calc.xlsx (a cópia recalculada pelo Excel).
Compara cada indicador da planilha com o mesmo indicador em dados_tratados/indicadores_2024.csv.
"""
import sys
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
arq = Path(sys.argv[1]) if len(sys.argv) > 1 else RAIZ / "planilhas" / "saude_fiscal_rj_2024_calc.xlsx"
ind = pd.read_csv(RAIZ / "dados_tratados" / "indicadores_2024.csv", dtype={"cod_ibge": str}).set_index("cod_ibge")

wb = load_workbook(arq, data_only=True)
ERROS = {"#DIV/0!", "#NAME?", "#VALUE!", "#REF!", "#N/A", "#NUM!", "#NULL!"}
print("Abas:", wb.sheetnames)

# 1) erros de fórmula em todas as abas
total_erros = 0
for ws in wb.worksheets:
    erros = [(c.coordinate, c.value) for linha in ws.iter_rows() for c in linha if isinstance(c.value, str) and c.value in ERROS]
    total_erros += len(erros)
    print(f"  {ws.title}: {ws.max_row} linhas x {ws.max_column} colunas | erros de fórmula: {len(erros)} {erros[:5]}")
print("Total de erros de fórmula:", total_erros)


def aba_como_tabela(nome):
    ws = wb[nome]
    cab = [c.value for c in ws[1]]
    linhas = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    return pd.DataFrame(linhas, columns=cab).dropna(how="all")


# 2) cada indicador da planilha contra o notebook
MAPA = {  # cabeçalho na planilha: (coluna do CSV, fator)
    "Saúde por habitante (R$)": ("saude_por_hab", 1), "Saúde, % da despesa": ("saude_pct_despesa", 100),
    "Receita por habitante (R$)": ("receita_por_hab", 1), "Despesa, % da receita": ("despesa_pct_receita", 100),
    "Pessoal total (R$)": ("pessoal_total", 1), "Pessoal total, % da RCL": ("pessoal_total_pct_rcl", 100),
    "Pessoal do Executivo, % da RCL": ("pessoal_pct_rcl", 100), "Restos a pagar, % da receita": ("rp_pct_receita", 100),
}
total_dif = 0
for aba in ("Estados", "Municípios RJ"):
    t = aba_como_tabela(aba)
    t = t.set_index("Código IBGE")
    for cab, (col, fator) in MAPA.items():
        a = pd.to_numeric(t[cab].replace("", np.nan), errors="coerce") * fator
        b = ind[col].reindex(t.index)
        if col == "pessoal_pct_rcl":     # o CSV traz o % arredondado que o ente declara; a planilha calcula pelo valor exato
            b = (ind.pessoal / ind.rcl_ajustada * 100).reindex(t.index)
        dif = (a - b).abs()
        tol = 1e-6 * b.abs().clip(lower=1) + 1e-6
        # onde só um dos lados é vazio, é diferença
        so_um = (a.isna() != b.isna())
        ruins = ((dif > tol) | so_um)
        if ruins.any():
            total_dif += int(ruins.sum())
            print(f"  DIFERENÇA em {aba} / {cab}: {int(ruins.sum())} entes, ex.: {list(t.index[ruins][:3])}")
    print(f"  {aba}: {len(t)} entes conferidos em {len(MAPA)} indicadores")
print("Total de diferenças planilha x notebook:", total_dif)

# 3) posições do RJ e medianas "com e sem"
est = aba_como_tabela("Estados").set_index("Código IBGE")
print("\nRJ (estado), posições na planilha:")
print(est.loc["33", [c for c in est.columns if str(c).startswith("Posição")]].to_dict())
cs = wb["Com e sem"]
print("\nCom e sem (linhas 5 a 8):")
for r in range(5, 9):
    print([cs.cell(row=r, column=c).value if not isinstance(cs.cell(row=r, column=c).value, float) else round(cs.cell(row=r, column=c).value, 4) for c in range(1, 7)])

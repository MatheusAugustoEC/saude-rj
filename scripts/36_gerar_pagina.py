"""Monta a página de entrega (pagina/index.html) e os arquivos para baixar (pagina/arquivos/).

Uso: python -X utf8 scripts/36_gerar_pagina.py   (depois de scripts/35_dados_pagina.py)
Lê pagina_fonte/ (molde, estilo e código), dados_tratados/pagina_dados.json e as malhas do IBGE; copia só arquivos que existem;
gera a planilha .xlsx (abas Sobre, Cidades, Regiões, Estados) e o .zip da Biblioteca. A apresentação .pptx é feita no navegador, com os mesmos dados.
"""
import json
import shutil
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
FONTE = RAIZ / "pagina_fonte"
SAIDA = RAIZ / "pagina"
ARQ = SAIDA / "arquivos"
BUNDLE = FONTE / "pptxgen.bundle.js"

dados = json.loads((RAIZ / "dados_tratados" / "pagina_dados.json").read_text(encoding="utf-8"))

# (grupo, caminho, título, descrição, para quem, tipo, etapa, ano)
CATALOGO = [
    ("Para ler", "documentos/01-hipoteses-e-resultados.pdf", "Hipóteses e resultados", "Eu apostei antes de olhar os números: o que previ, o que deu e onde errei.", "Todos", "PDF", "Controlar", "2023 a 2025"),
    ("Para ler", "documentos/02-plano-de-controle.pdf", "Plano de controle", "O que acompanhar daqui para frente, com que frequência e quem age.", "Gestão", "PDF", "Controlar", "2023 a 2025"),
    ("Para ler", "documentos/03-kaizen-do-projeto.pdf", "Kaizen do projeto", "O que mais atrapalhou, as três mudanças e o que deu certo.", "Todos", "PDF", "Controlar", "2023 a 2025"),
    ("Para ler", "fases/fase-2a-measure-qualidade/relatorio.pdf", "Relatório do Measure: qualidade dos dados", "Como os dados foram conferidos e tratados.", "Analista", "PDF", "Medir", "2024"),
    ("Para ler", "fases/fase-2b-measure-baseline/relatorio.pdf", "Relatório do Measure: ponto de partida", "O gasto por habitante, medido, com os limites do dado.", "Analista", "PDF", "Medir", "2024"),
    ("Para ler", "fases/fase-3-analyze/relatorio.pdf", "Relatório do Analyze", "O que explica a diferença entre as cidades e o que não explica.", "Analista", "PDF", "Analisar", "2023 e 2024"),
    ("Para ler", "fases/fase-4-improve/relatorio.pdf", "Relatório do Improve", "Sete ideias testadas, cada uma com veredito.", "Gestão", "PDF", "Melhorar", "2023 e 2024"),
    ("Para ler", "fases/fase-5-control/relatorio.pdf", "Relatório do Control", "As conclusões testadas com os dados de 2025.", "Controle", "PDF", "Controlar", "2025"),
    ("Imagens e gráficos", "imagens/sipoc.png", "SIPOC", "Fornecedores, entradas, processo, saídas e clientes do projeto.", "Todos", "Imagem", "Definir", "2023 a 2025"),
    ("Imagens e gráficos", "imagens/mapa_processo.png", "Mapa do processo", "Da coleta dos dados à entrega.", "Todos", "Imagem", "Definir", "2023 a 2025"),
    ("Imagens e gráficos", "imagens/marco2_measure.png", "Resumo do Measure", "A situação medida, em uma imagem.", "Todos", "Imagem", "Medir", "2024"),
    ("Imagens e gráficos", "imagens/marco2_mapa_regioes.png", "Mapa das regiões", "Gasto em saúde por habitante das regiões do estado.", "Todos", "Imagem", "Medir", "2024"),
    ("Imagens e gráficos", "imagens/marco2_mapa_C_regioes.png", "Mapa das regiões, com tabela", "As regiões e os valores lado a lado.", "Todos", "Imagem", "Medir", "2024"),
    ("Imagens e gráficos", "documentos/fonte/img/veredictos.png", "Veredito das nove conclusões", "O que se manteve e o que enfraqueceu com os dados de 2025.", "Todos", "Imagem", "Controlar", "2025"),
    ("Imagens e gráficos", "documentos/fonte/img/chegada_dca_2025.png", "Chegada do balanço anual de 2025", "Quantas cidades entregaram em cada mês.", "Controle", "Imagem", "Controlar", "2025"),
    ("Planilhas, dados e apresentação", "indicadores_2024.csv", "Indicadores de 2024", "Gasto, receita e despesa por habitante dos 92 municípios e 27 estados.", "Analista", "Dados", "Medir", "2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/analyze_a2_residuos.csv", "Resíduos do modelo", "Quanto cada cidade gasta a mais ou a menos do que a receita indica.", "Analista", "Dados", "Analisar", "2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/improve_r1_triagem.csv", "Triagem de auditoria", "Cidades fora do intervalo de previsão, 2023 e 2024.", "Controle", "Dados", "Melhorar", "2023 e 2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/improve_r2_pares.csv", "Cidades parecidas", "Faixa de receita de cada cidade e desvio contra região e contra pares.", "Gestão", "Dados", "Melhorar", "2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/improve_r4_estresse.csv", "Teste de estresse dos repasses", "Cenários de queda dos repasses, com os três pressupostos.", "Gestão", "Dados", "Melhorar", "2023 a 2025"),
    ("Planilhas, dados e apresentação", "dados_tratados/improve_r5_r6_siops.csv", "Folga e conferência com o SIOPS", "Folga sobre o mínimo de 15% e razão entre SIOPS e Siconfi.", "Controle", "Dados", "Melhorar", "2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/improve_r7_limites.csv", "Limites da LRF", "Pessoal, dívida e restos a pagar declarados.", "Controle", "Dados", "Melhorar", "2024"),
    ("Planilhas, dados e apresentação", "dados_tratados/control_2025_resultados.csv", "Veredito de 2025", "As nove conclusões testadas com os dados de 2025.", "Todos", "Dados", "Controlar", "2025"),
    ("Planilhas, dados e apresentação", "dados_tratados/control_2025_residuos.csv", "Triagem de 2025", "Resíduo de cada cidade com o balanço de 2025.", "Controle", "Dados", "Controlar", "2025"),
    ("Planilhas, dados e apresentação", "dados_tratados/analyze_ibge_indicadores.csv", "Indicadores do IBGE", "Idade e saneamento por cidade, Censo 2022.", "Analista", "Dados", "Analisar", "2022"),
    ("Planilhas, dados e apresentação", "pagina/arquivos/saude-fiscal-rj-cidades.xlsx", "Planilha das cidades", "Todas as cidades, regiões e estados desta página, em uma planilha com aba Sobre.", "Todos", "Planilha", "Medir", "2024"),
]

# explicações curtas de cada item da Biblioteca (o que é e por que existe)
DESC = {
    "SIPOC": "Antes de olhar qualquer número, desenhei quem fornece os dados, o que entra, o que o projeto faz, o que sai e quem usa. É o ponto de partida: mostra de onde cada informação vem e para quem ela serve.",
    "Mapa do processo": "O caminho que cada número percorre, da coleta no Siconfi à conferência, ao modelo e a esta página. Mostra onde o dado poderia se perder e onde fiz a conferência.",
    "Resumo do Measure": "Antes de explicar o número, desconfiei dele. A imagem conta a conferência: 118 de 119 entes batiam com a prestação de contas anual (119 de 119 depois da correção), a população do Siconfi foi trocada pela estimativa do IBGE, e o maior e o menor gasto por habitante trocaram de município. O ponto de partida ficou em R$ 2.319 por habitante, a mediana das 92 cidades.",
    "Mapa das regiões": "O estado do Rio dividido em regiões, com o gasto em saúde por habitante de cada uma. A Costa Verde, com 3 cidades, tem o maior valor; a capital, o menor.",
    "Mapa das regiões, com tabela": "O mesmo mapa com a tabela ao lado: cidades, população, gasto e receita por habitante de cada região. As regiões com poucas cidades (marcadas com †) só descrevem; não valem como tendência.",
    "Veredito das nove conclusões": "Escrevi o que esperava antes de abrir os dados de 2025. Aqui está o resultado: oito conclusões se mantiveram e uma, a lista de cidades da triagem, enfraqueceu.",
    "Chegada do balanço anual de 2025": "O balanço anual de cada prefeitura leva meses para chegar ao Tesouro. A curva mostra quantas já tinham entregado a cada mês e explica por que a triagem de 2025 só fecha depois.",
    "Indicadores de 2024": "A base de tudo: para cada município e estado, o gasto em saúde, a receita, a despesa e a população. As colunas principais são saude_por_hab, receita_por_hab e saude_pct_despesa.",
    "Resíduos do modelo": "Responde: dado o que a cidade arrecada, quanto ela deveria gastar em saúde? A coluna residuo_pct mostra a distância, em %, entre o gasto declarado e o previsto.",
    "Triagem de auditoria": "A lista de cidades fora do intervalo esperado em 2023 e 2024. res_pct é a distância do previsto, e fora_ic diz se o gasto ficou abaixo do intervalo de 95%.",
    "Cidades parecidas": "A faixa de receita de cada cidade (de 0 a 4) e quanto ela gasta em relação à mediana da região (dev_reg) e à mediana das cidades da mesma faixa (dev_par).",
    "Teste de estresse dos repasses": "E se os repasses caírem? Para quedas de 10%, 20% e 30%, três cenários (otimista, provável e pessimista) de quanto cairia o gasto em saúde.",
    "Folga e conferência com o SIOPS": "Quanto cada cidade aplica acima do mínimo de 15% (folga_pp), a queda que zeraria essa folga (queda_zera_pct) e a razão entre o SIOPS e o Siconfi (razao).",
    "Limites da LRF": "O que cada município declara de pessoal, dívida e restos a pagar, contra os limites da Lei de Responsabilidade Fiscal.",
    "Veredito de 2025": "As nove conclusões e como cada uma resistiu aos dados de 2025: se mantém ou enfraquece, com e sem as declarações entregues por último.",
    "Triagem de 2025": "A distância do previsto de cada cidade com os dados de 2025 (res_pct) e se a declaração é recente (recente), ou seja, ainda sujeita a mudança.",
    "Indicadores do IBGE": "Idade e saneamento de cada cidade no Censo 2022, usados para testar se explicam o gasto. Não explicam.",
    "Planilha das cidades": "A mesma informação desta página em uma planilha: aba Cidades (92 linhas), Regiões (8) e Estados (27), mais uma aba Sobre com a fonte e os cuidados. Para outro recorte, use o filtro da página Gestor.",
}


def tam(p):
    n = p.stat().st_size
    return f"{n / 1024:.0f} KB" if n < 1024 * 1024 else f"{n / 1024 / 1024:.1f} MB"


def planilha(dst):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sobre"
    for i, t in enumerate([
        "Saúde fiscal do Rio de Janeiro: planilha da página de entrega",
        f"Versão {dados['meta']['versao']}, {dados['meta']['data']}. Janela principal: {dados['meta']['janela_principal']}.",
        "Fonte: Siconfi (Tesouro Nacional), SIOPS (Ministério da Saúde) e IBGE. Valores declarados pelos entes, não auditados.",
        "Gasto em saúde por habitante não mede qualidade nem necessidade. O estado (esfera estadual) não é comparável diretamente com os municípios.",
        "Aba Cidades: 92 municípios. Aba Regiões: 8 regiões de governo (soma do gasto ÷ soma da população). Aba Estados: 27 estados, 2024.",
        f"Repositório: {dados['meta']['repositorio']}",
    ], 1):
        ws.cell(i, 1, t).font = Font(bold=(i == 1), size=13 if i == 1 else 11)
    ws.column_dimensions["A"].width = 130

    def aba(nome, cab, linhas, larg):
        w = wb.create_sheet(nome)
        w.append([c[0] for c in cab])
        for l in linhas:
            w.append([l.get(c[1]) for c in cab])
        for j, c in enumerate(cab, 1):
            cel = w.cell(1, j)
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="002147")
            cel.alignment = Alignment(wrap_text=True, vertical="center")
            w.column_dimensions[get_column_letter(j)].width = larg if j > 1 else 30
            for r in range(2, len(linhas) + 2):
                if c[2]:
                    w.cell(r, j).number_format = c[2]
        w.freeze_panes = "B2"
        w.auto_filter.ref = w.dimensions
        w.row_dimensions[1].height = 45

    aba("Cidades", [("Cidade", "n", None), ("Região", "r", None), ("Grupo do estudo", "g", None), ("População", "pop", "#,##0"), ("Gasto em saúde por habitante (R$, 2024)", "gh", "#,##0.0"), ("Receita por habitante (R$, 2024)", "rh", "#,##0.0"),
                         ("Parte da despesa em saúde (%)", "pd", "0.00"), ("Saúde como parte da receita (%)", "fa", "0.00"), ("Aplicado em saúde, SIOPS (%)", "ap", "0.00"), ("Folga sobre o mínimo de 15% (p.p.)", "fo", "0.00"),
                         ("Queda do valor aplicado que zera a folga (%)", "qz", "0.0"), ("Parte da receita que vem de repasses (%)", "pt", "0.0"), ("Distância do previsto, 2023 (%)", "r23", "0.0"), ("Distância do previsto, 2024 (%)", "r24", "0.0"),
                         ("Distância do previsto, 2025 (%)", "r25", "0.0"), ("Pessoal (% da RCL, 2024)", "pes", "0.00"), ("Limite prudencial (%)", "prud", "0.00")], dados["cidades"], 16)
    aba("Regiões", [("Região", "k", None), ("Cidades", "n", "0"), ("População", "pop", "#,##0"), ("Gasto por habitante (R$)", "gh", "#,##0.0"), ("Receita por habitante (R$)", "rh", "#,##0.0"), ("Parte da despesa em saúde (%)", "pd", "0.00"), ("Repasses na receita (%)", "tr", "0.0")], dados["regioes"], 16)
    aba("Estados", [("Estado", "n", None), ("Sigla", "s", None), ("Gasto por habitante (R$)", "gh", "#,##0.0"), ("Parte da despesa em saúde (%)", "pd", "0.00"), ("Receita por habitante (R$)", "rh", "#,##0.0")], dados["ufs"], 16)
    wb.save(dst)


def main():
    if SAIDA.exists():
        shutil.rmtree(SAIDA)
    ARQ.mkdir(parents=True)
    planilha(ARQ / "saude-fiscal-rj-cidades.xlsx")
    bib, zipar = [], []
    for g, rel, titulo, desc, para, tipo, etapa, ano in CATALOGO:
        src = RAIZ / (rel if "/" in rel else f"dados_tratados/{rel}")
        if not src.exists():
            print("  ausente, fora da Biblioteca:", rel)
            continue
        nome = f"{src.parent.name}-relatorio.pdf" if src.name == "relatorio.pdf" else src.name
        dst = ARQ / nome
        if src.resolve() != dst.resolve():
            shutil.copy2(src, dst)
        item = {"grupo": g, "titulo": titulo, "desc": DESC.get(titulo, desc), "para": para, "tipo": tipo, "etapa": etapa, "ano": ano, "arq": f"arquivos/{nome}", "tam": tam(dst), "data": datetime.fromtimestamp(src.stat().st_mtime).strftime("%d/%m/%Y"), "img": tipo == "Imagem", "alt": titulo}
        bib.append(item)
        zipar.append(dst)
    for extra in ("imagens/sipoc.png", "imagens/mapa_processo.png"):
        d = ARQ / Path(extra).name
        if not d.exists():
            shutil.copy2(RAIZ / extra, d)
    bib.append({"grupo": "Para reproduzir e citar", "titulo": "Repositório no GitHub", "desc": "Código, dados e relatórios, com instruções para refazer tudo.", "para": "Analista", "tipo": "Código", "etapa": "Controlar", "ano": "2023 a 2025", "arq": dados["meta"]["repositorio"], "tam": "repositório", "data": dados["meta"]["data"], "ext": True})
    req = RAIZ / "requirements.txt"
    if req.exists():
        shutil.copy2(req, ARQ / "requirements.txt")
        bib.append({"grupo": "Para reproduzir e citar", "titulo": "Bibliotecas usadas", "desc": "Lista das bibliotecas Python para refazer as contas.", "para": "Analista", "tipo": "Código", "etapa": "Controlar", "ano": "2023 a 2025", "arq": "arquivos/requirements.txt", "tam": tam(req), "data": datetime.fromtimestamp(req.stat().st_mtime).strftime("%d/%m/%Y")})
        zipar.append(ARQ / "requirements.txt")
    with zipfile.ZipFile(ARQ / "saude-fiscal-rj-biblioteca.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for f in zipar:
            z.write(f, f.name)
    geo_uf = (RAIZ / "dados_brutos/malha_ibge_api/br_uf_qualidade_minima.json").read_text(encoding="utf-8")
    geo_rj = (RAIZ / "dados_brutos/malha_ibge_api/rj_municipios_qualidade_minima.json").read_text(encoding="utf-8")
    js = "\n".join((FONTE / n).read_text(encoding="utf-8") for n in ("app.js", "util.js", "views1.js", "views2.js", "views3.js"))
    html = (FONTE / "template.html").read_text(encoding="utf-8")
    for k, v in (("/*CSS*/", (FONTE / "estilo.css").read_text(encoding="utf-8")), ("/*DADOS*/null", json.dumps(dados, ensure_ascii=False, separators=(",", ":"))), ("/*GEO_UF*/null", json.dumps(json.loads(geo_uf), separators=(",", ":"))),
                 ("/*GEO_RJ*/null", json.dumps(json.loads(geo_rj), separators=(",", ":"))), ("/*BIB*/null", json.dumps(bib, ensure_ascii=False, separators=(",", ":"))), ("/*JS*/", js)):
        assert k in html, k
        html = html.replace(k, v, 1)
    (SAIDA / "index.html").write_text(html, encoding="utf-8")
    shutil.copy2(BUNDLE, SAIDA / "pptxgen.bundle.js")
    print(f"Gerado pagina/index.html ({(SAIDA / 'index.html').stat().st_size / 1024:.0f} KB); {len(bib)} itens na Biblioteca; zip com {len(zipar)} arquivos.")


if __name__ == "__main__":
    main()

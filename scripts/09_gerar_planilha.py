"""Gera a planilha Excel do projeto: dados, indicadores (com fórmulas), posições, "com e sem", gráficos e legenda.

Uso: python -u -I scripts/09_gerar_planilha.py
Lê dados_tratados/indicadores_2024.csv (gerado pelo notebook da EDA) e dados_brutos/siconfi_rgf_2024.csv
(só para separar a despesa de pessoal por Poder). Grava planilhas/saude_fiscal_rj_2024.xlsx.

Regras da planilha:
- Texto AZUL = dado que veio da API do Siconfi (entrada). Texto PRETO = fórmula, que o leitor pode conferir.
- Os indicadores são fórmulas sobre as colunas de entrada, para qualquer pessoa conferir o cálculo.
Depois de gerar, rodar a recalculagem (recalc.py da skill de planilhas) para gravar os valores calculados.
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as LETRA

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "planilhas" / "saude_fiscal_rj_2024.xlsx"
COD_RJ, COD_CAPITAL = "33", "3304557"

FONTE = "Arial"
F_NORMAL = Font(name=FONTE, size=10)
F_NEGRITO = Font(name=FONTE, size=10, bold=True)
F_ENTRADA = Font(name=FONTE, size=10, color="0000FF")          # dado da API
F_CAB = Font(name=FONTE, size=10, bold=True)
F_TITULO = Font(name=FONTE, size=16, bold=True)
F_SUB = Font(name=FONTE, size=12, bold=True)
FILL_CAB = PatternFill("solid", fgColor="E1E0D9")
FILL_RJ = PatternFill("solid", fgColor="DCEBFB")
FILL_ESPECIAL = PatternFill("solid", fgColor="FFF2CC")
BORDA = Border(bottom=Side(style="thin", color="C3C2B7"))
FMT = {"int": "#,##0", "rs": '"R$" #,##0', "pct": "0.0%", "dec": "0.00", "txt": "@"}

# ------------------------------------------------------------------------------------------ dados
ind = pd.read_csv(RAIZ / "dados_tratados" / "indicadores_2024.csv", dtype={"cod_ibge": str})
ind["flag_especial"] = ind.flag_especial.fillna("")      # texto vazio é lido do CSV como "sem valor"
rgf = pd.read_csv(RAIZ / "dados_brutos" / "siconfi_rgf_2024.csv", dtype={"cod_ibge": str})
a1 = rgf[(rgf.anexo == "RGF-Anexo 01") & (rgf.cod_conta == "DespesaComPessoalTotal") & (rgf.coluna == "Valor")]
poder = a1.groupby(["cod_ibge", "co_poder"]).valor.first().unstack()
for p in "ELJMD":
    if p not in poder.columns:
        poder[p] = np.nan

# Conferência: a separação por Poder tem de reproduzir os totais que o notebook já calculou
chk = ind.set_index("cod_ibge")
assert np.allclose(poder["E"].reindex(chk.index), chk.pessoal, equal_nan=True), "pessoal do Executivo não bate com o notebook"
assert np.allclose(poder.sum(axis=1).reindex(chk.index), chk.pessoal_total, equal_nan=True), "pessoal total não bate com o notebook"


def vazio(v):
    return None if pd.isna(v) else v


def observacao(r):
    partes = []
    if r.flag_especial:
        partes.append(r.flag_especial.capitalize())
    if r.relatorio_simplificado:
        partes.append("Relatório simplificado")
    if r.poderes_incompletos:
        partes.append("Poderes incompletos no RGF: a soma de pessoal fica subestimada")
    if r.limite_dcl_fora_padrao:
        partes.append("Limite de dívida declarado fora do padrão (a checar)")
    if r.limite_declarado_como_percentual:
        partes.append("Limite de dívida declarado já como percentual")
    if pd.isna(r.limite_senado):
        partes.append("Sem limite de dívida declarado")
    if pd.isna(r.rp_saldo):
        partes.append("Sem linha de restos a pagar")
    return "; ".join(partes)


# --------------------------------------------------------------------------- especificação das colunas
# Cada coluna: (chave, cabeçalho, tipo, formato, largura, descrição, origem)
# tipo: "entrada" (azul) ou "formula" (preto, com modelo usando {chave} para a célula da mesma linha e {chave_c} para a letra da coluna)
def colunas_comuns(poderes):
    pod = [(f"p{c}", f"Pessoal {n} (R$)", "entrada", "rs", 16, f"Despesa total com pessoal do {n}",
            "RGF Anexo 01, conta DespesaComPessoalTotal, coluna 'Valor'") for c, n in poderes]
    cols = [
        ("nome", "Ente", "entrada", "txt", 28, "Nome do ente", "Cadastro do Siconfi (`/entes`)"),
        ("cod", "Código IBGE", "entrada", "txt", 11, "Código IBGE (2 dígitos para estado, 7 para município)", "Cadastro do Siconfi"),
        ("pop", "População", "entrada", "int", 12, "População do ente, estimativa 2024 do IBGE", "IBGE, SIDRA, tabela 6579, variável 9324"),
        ("receita", "Receita realizada (R$)", "entrada", "rs", 17, "Receita realizada no ano, sem intra-orçamentárias",
         "RREO Anexo 01, ReceitasExcetoIntraOrcamentarias, 'Até o Bimestre (c)'"),
        ("despesa", "Despesa empenhada, Anexo 01 (R$)", "entrada", "rs", 17, "Despesa empenhada no ano, sem intra-orçamentárias",
         "RREO Anexo 01, DespesasExcetoIntraOrcamentarias, 'Despesas empenhadas até o bimestre (f)'"),
        ("despesa02", "Despesa total, Anexo 02 (R$)", "entrada", "rs", 17, "Total de despesa do Anexo 02, sem intra-orçamentárias (em Armação dos Búzios, o Anexo 02 é incoerente e vale o total do Anexo 01)",
         "RREO Anexo 02, linha 'DESPESAS (EXCETO INTRA-ORÇAMENTÁRIAS) (I)'; Búzios: Anexo 01"),
        ("saude", "Despesa em Saúde (R$)", "entrada", "rs", 17, "Despesa empenhada na função Saúde, sem intra-orçamentárias",
         "RREO Anexo 02, linha 'Saúde', 'Despesas empenhadas até o bimestre (b)'; Armação dos Búzios: DCA Anexo I-E, função 10 (checagem A2)"),
        ("rcl", "RCL ajustada (R$)", "entrada", "rs", 17, "Receita corrente líquida ajustada, base dos limites de pessoal",
         "RGF Anexo 01, ReceitaCorrenteLiquidaAjustada, 'Valor'"),
    ] + pod + [
        ("dcl", "Dívida líquida, % da RCL (declarado)", "entrada", "pct", 16, "Dívida consolidada líquida como % da RCL ajustada para endividamento",
         "RGF Anexo 02, PercentualDaDCLSobreARCL, coluna do último período"),
        ("lim", "Limite de dívida do Senado, % da RCL (declarado)", "entrada", "pct", 18,
         "Limite de endividamento declarado, como % da RCL (limite ÷ RCL de endividamento, ou o próprio percentual quando o ente o declarou assim)",
         "RGF Anexo 02, LimiteDefinidoPorResolucaoDoSenadoFederal"),
        ("rp", "Restos a pagar, saldo (R$)", "entrada", "rs", 16, "Saldo total de restos a pagar, sem intra-orçamentários",
         "RREO Anexo 07, SaldoTotal, linha 'RESTOS A PAGAR (EXCETO INTRA-ORÇAMENTÁRIOS) (I)'"),
    ]
    return cols


def colunas_formulas(primeiro_p, ultimo_p, minimo_poderes):
    return [
        ("saude_hab", "Saúde por habitante (R$)", "formula", "rs", 14, "Saúde ÷ população", "Fórmula: Despesa em Saúde ÷ População",
         'IFERROR({saude}/{pop},"")'),
        ("saude_pct", "Saúde, % da despesa", "formula", "pct", 12, "Fatia da despesa total (Anexo 02) que vai para Saúde", "Fórmula: Despesa em Saúde ÷ Despesa total do Anexo 02",
         'IFERROR({saude}/{despesa02},"")'),
        ("rec_hab", "Receita por habitante (R$)", "formula", "rs", 14, "Receita realizada ÷ população", "Fórmula: Receita realizada ÷ População",
         'IFERROR({receita}/{pop},"")'),
        ("desp_pct", "Despesa, % da receita", "formula", "pct", 12, "Despesa empenhada ÷ receita realizada", "Fórmula: Despesa empenhada ÷ Receita realizada",
         'IFERROR({despesa}/{receita},"")'),
        ("pess_tot", "Pessoal total (R$)", "formula", "rs", 16, "Soma do pessoal de todos os Poderes", "Fórmula: soma das colunas de pessoal por Poder",
         "SUM({" + primeiro_p + "}:{" + ultimo_p + "})"),
        ("n_pod", "Poderes com dado", "formula", "int", 9, "Quantos Poderes têm dado no RGF", "Fórmula: contagem das colunas de pessoal preenchidas",
         "COUNT({" + primeiro_p + "}:{" + ultimo_p + "})"),
        ("pess_pct", "Pessoal total, % da RCL", "formula", "pct", 12, "Indicador principal de pessoal: soma dos Poderes ÷ RCL ajustada",
         "Fórmula: Pessoal total ÷ RCL ajustada", 'IFERROR({pess_tot}/{rcl},"")'),
        ("pess_exec", "Pessoal do Executivo, % da RCL", "formula", "pct", 12, "Pessoal do Executivo ÷ RCL ajustada", "Fórmula: Pessoal Executivo ÷ RCL ajustada",
         'IFERROR({' + primeiro_p + '}/{rcl},"")'),
        ("rp_pct", "Restos a pagar, % da receita", "formula", "pct", 12, "Saldo de restos a pagar ÷ receita realizada", "Fórmula: Restos a pagar ÷ Receita realizada",
         'IF({rp}="","",IFERROR({rp}/{receita},""))'),
    ]


def colunas_posicao(minimo_poderes, com_divida):
    cols = [
        ("pos_saude_hab", "Posição: saúde por habitante", "formula", "int", 11, "Posição entre os entes do grupo (1 = maior valor)", "Fórmula: RANK da coluna Saúde por habitante",
         'IFERROR(RANK({saude_hab},{saude_hab_c}$2:{saude_hab_c}${ult}),"")'),
        ("pos_saude_pct", "Posição: saúde % da despesa", "formula", "int", 11, "Posição entre os entes do grupo (1 = maior valor)", "Fórmula: RANK da coluna Saúde, % da despesa",
         'IFERROR(RANK({saude_pct},{saude_pct_c}$2:{saude_pct_c}${ult}),"")'),
        ("pos_desp", "Posição: despesa % da receita", "formula", "int", 11, "Posição entre os entes do grupo (1 = maior valor)", "Fórmula: RANK da coluna Despesa, % da receita",
         'IFERROR(RANK({desp_pct},{desp_pct_c}$2:{desp_pct_c}${ult}),"")'),
        ("pos_pess", "Posição: pessoal total % da RCL", "formula", "int", 11,
         "Posição entre os entes com Poderes completos (1 = maior); 'n/c' = não comparável",
         "Fórmula: conta quantos entes completos têm valor maior",
         'IF({n_pod}>=' + str(minimo_poderes) + ',COUNTIFS({n_pod_c}$2:{n_pod_c}${ult},">=' + str(minimo_poderes)
         + '",{pess_pct_c}$2:{pess_pct_c}${ult},">"&{pess_pct})+1,"n/c")'),
    ]
    if com_divida:
        cols.append(("pos_dcl", "Posição: dívida % da RCL", "formula", "int", 11, "Posição entre os entes do grupo (1 = maior valor)",
                     "Fórmula: RANK da coluna Dívida líquida, % da RCL", 'IFERROR(RANK({dcl},{dcl_c}$2:{dcl_c}${ult}),"")'))
    return cols


# --------------------------------------------------------------------------------------- escrever aba
def escrever_aba(ws, tabela, colunas, destaque_cod=None):
    letras = {c[0]: LETRA(i + 1) for i, c in enumerate(colunas)}
    ult = len(tabela) + 1
    for i, c in enumerate(colunas, start=1):
        cel = ws.cell(row=1, column=i, value=c[1])
        cel.font, cel.fill = F_CAB, FILL_CAB
        cel.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[LETRA(i)].width = c[4]
    ws.row_dimensions[1].height = 62
    for r, (_, linha) in enumerate(tabela.iterrows(), start=2):
        refs = {k: f"{v}{r}" for k, v in letras.items()}
        refs.update({f"{k}_c": v for k, v in letras.items()})
        refs["ult"] = ult
        for i, c in enumerate(colunas, start=1):
            chave, _, tipo, fmt = c[0], c[1], c[2], c[3]
            if tipo == "entrada":
                valor = linha[chave]
                valor = None if (valor is None or (not isinstance(valor, str) and pd.isna(valor))) else valor
                cel = ws.cell(row=r, column=i, value=valor)
                cel.font = F_NORMAL if chave in ("nome", "obs", "simplificado", "caso") else F_ENTRADA
            else:
                cel = ws.cell(row=r, column=i, value="=" + c[7].format(**refs))
                cel.font = F_NORMAL
            cel.number_format = FMT[fmt]
            cel.border = BORDA
            if fmt == "txt":
                cel.alignment = Alignment(vertical="center")
        if destaque_cod is not None and linha["cod"] == destaque_cod:
            for i in range(1, len(colunas) + 1):
                ws.cell(row=r, column=i).fill = FILL_RJ
                ws.cell(row=r, column=i).font = Font(name=FONTE, size=10, bold=True, color=ws.cell(row=r, column=i).font.color)
        elif linha.get("especial", False):
            for i in range(1, len(colunas) + 1):
                ws.cell(row=r, column=i).fill = FILL_ESPECIAL
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{LETRA(len(colunas))}{ult}"
    return letras, ult


def preparar(df, grupo, poderes):
    d = df[df.grupo == grupo].copy().sort_values("nome")
    out = pd.DataFrame({
        "nome": d.nome.values, "cod": d.cod_ibge.values, "pop": d.populacao.values, "receita": d.receita.values,
        "despesa": d.despesa.values, "despesa02": d.despesa_base_saude.values, "saude": d.saude.values, "rcl": d.rcl_ajustada.values,
        "dcl": (d.dcl_pct_rcl / 100).values, "lim": (d.limite_dcl_pct / 100).values, "rp": d.rp_saldo.values,
        "obs": [observacao(r) for r in d.itertuples()], "especial": (d.flag_especial != "").values,
        "simplificado": np.where(d.relatorio_simplificado, "sim", "não"), "caso": d.flag_especial.values,
    })
    for c, _ in poderes:
        out[f"p{c}"] = poder[c].reindex(d.cod_ibge).values
    return out


wb = Workbook()
ws_leia = wb.active
ws_leia.title = "Leia-me"

# ------------------------------------------------------------------------------------- aba Estados
PODERES_EST = [("E", "Executivo"), ("L", "Legislativo"), ("J", "Judiciário"), ("M", "Ministério Público"), ("D", "Defensoria")]
tab_est = preparar(ind, "Estado", PODERES_EST)
cols_est = (colunas_comuns(PODERES_EST) + colunas_formulas("pE", "pD", 4) + colunas_posicao(4, True)
            + [("obs", "Observações", "entrada", "txt", 60, "Avisos sobre o ente (casos especiais, dados ausentes, comparabilidade)", "Notebook da EDA")])
ws_est = wb.create_sheet("Estados")
let_est, ult_est = escrever_aba(ws_est, tab_est, cols_est, destaque_cod=COD_RJ)

# --------------------------------------------------------------------------------- aba Municípios RJ
PODERES_MUN = [("E", "Executivo"), ("L", "Legislativo")]
tab_mun = preparar(ind, "Município do RJ", PODERES_MUN)
cols_mun = (colunas_comuns(PODERES_MUN)
            + [("simplificado", "Relatório simplificado?", "entrada", "txt", 12, "O município entrega RREO e RGF simplificados", "Registro da coleta (extrato de entregas)"),
               ("caso", "Caso especial", "entrada", "txt", 26, "Dado atípico tratado 'com e sem'", "Notebook da EDA")]
            + colunas_formulas("pE", "pL", 2) + colunas_posicao(2, False)
            + [("obs", "Observações", "entrada", "txt", 60, "Avisos sobre o ente", "Notebook da EDA")])
ws_mun = wb.create_sheet("Municípios RJ")
let_mun, ult_mun = escrever_aba(ws_mun, tab_mun, cols_mun, destaque_cod=COD_CAPITAL)

# ------------------------------------------------------------------------------------- aba Com e sem
ws_cs = wb.create_sheet("Com e sem")
ws_cs["A1"] = "Resultados com e sem os casos especiais (municípios do RJ)"
ws_cs["A1"].font = F_SUB
ws_cs["A2"] = ("Mediana de cada indicador em quatro versões: todos os municípios; sem os 4 de dado atípico; sem os 12 de relatório simplificado; "
               "sem os dois grupos. Tudo é fórmula sobre a aba 'Municípios RJ'. Os 4 atípicos são os marcados em amarelo naquela aba.")
ws_cs["A2"].font, ws_cs["A2"].alignment = F_NORMAL, Alignment(wrap_text=True, vertical="top")
ws_cs.merge_cells("A2:G2")
ws_cs.row_dimensions[2].height = 42
IND_CS = [("saude_hab", "Saúde por habitante (mediana)", "rs"), ("saude_pct", "Saúde, % da despesa (mediana)", "pct"),
          ("pess_pct", "Pessoal total, % da RCL (mediana)", "pct"), ("desp_pct", "Despesa, % da receita (mediana)", "pct")]
cab = ["Versão", "nº de municípios"] + [x[1] for x in IND_CS]
for i, h in enumerate(cab, start=1):
    c = ws_cs.cell(row=4, column=i, value=h)
    c.font, c.fill = F_CAB, FILL_CAB
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws_cs.row_dimensions[4].height = 44
ws_cs.column_dimensions["A"].width = 38
for i in range(2, 7):
    ws_cs.column_dimensions[LETRA(i)].width = 18

N = len(tab_mun)
INI_AUX = 12                                    # primeira linha do bloco auxiliar
AUX_FIM = INI_AUX + N - 1
VERSOES = [("Todos os municípios", None), ("Sem os 4 de dado atípico", "atip"), ("Sem os 12 simplificados", "simp"), ("Sem os atípicos e sem os simplificados", "ambos")]
# bloco auxiliar: cabeçalhos
ws_cs.cell(row=INI_AUX - 2, column=1, value="Bloco auxiliar (alimenta as medianas acima): cada coluna repete o indicador só para os municípios que ficam na versão.").font = F_NEGRITO
aux_cab = ["Município", "Caso especial", "Simplificado?"]
for k, (chave, rotulo, _) in enumerate(IND_CS):
    for v in ("atip", "simp", "ambos"):
        aux_cab.append(f"{rotulo.split(' (')[0]} | {dict(atip='sem atípicos', simp='sem simplificados', ambos='sem ambos')[v]}")
for i, h in enumerate(aux_cab, start=1):
    c = ws_cs.cell(row=INI_AUX - 1, column=i, value=h)
    c.font, c.fill = F_CAB, FILL_CAB
    c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
ws_cs.row_dimensions[INI_AUX - 1].height = 60
M = "'Municípios RJ'!"
for k in range(N):
    r_aux, r_src = INI_AUX + k, 2 + k
    ws_cs.cell(row=r_aux, column=1, value=f"={M}{let_mun['nome']}{r_src}").font = F_NORMAL
    ws_cs.cell(row=r_aux, column=2, value=f'={M}{let_mun["caso"]}{r_src}&""').font = F_NORMAL
    ws_cs.cell(row=r_aux, column=3, value=f"={M}{let_mun['simplificado']}{r_src}").font = F_NORMAL
    col = 4
    for chave, _, fmt in IND_CS:
        src = f"{M}{let_mun[chave]}{r_src}"
        conds = {"atip": f'$B{r_aux}=""', "simp": f'$C{r_aux}="não"', "ambos": f'AND($B{r_aux}="",$C{r_aux}="não")'}
        for v in ("atip", "simp", "ambos"):
            c = ws_cs.cell(row=r_aux, column=col, value=f'=IF({conds[v]},{src},"")')
            c.font, c.number_format = F_NORMAL, FMT[fmt]
            col += 1
for r_v, (nome_v, v) in enumerate(VERSOES, start=5):
    ws_cs.cell(row=r_v, column=1, value=nome_v).font = F_NEGRITO if v is None else F_NORMAL
    col_aux = 4
    for j, (chave, _, fmt) in enumerate(IND_CS):
        if v is None:
            rng = f"{M}{let_mun[chave]}2:{let_mun[chave]}{ult_mun}"
        else:
            idx = 4 + j * 3 + ("atip", "simp", "ambos").index(v)
            rng = f"{LETRA(idx)}{INI_AUX}:{LETRA(idx)}{AUX_FIM}"
        c = ws_cs.cell(row=r_v, column=3 + j, value=f"=MEDIAN({rng})")
        c.font, c.number_format, c.border = F_NORMAL, FMT[fmt], BORDA
        if j == 0:
            idx0 = 4 + ("atip", "simp", "ambos").index(v) if v else None
            n = (f"=COUNT({M}{let_mun['saude_hab']}2:{let_mun['saude_hab']}{ult_mun})" if v is None
                 else f"=COUNT({LETRA(idx0)}{INI_AUX}:{LETRA(idx0)}{AUX_FIM})")
            cn = ws_cs.cell(row=r_v, column=2, value=n)
            cn.font, cn.number_format, cn.border = F_NORMAL, FMT["int"], BORDA
    ws_cs.cell(row=r_v, column=1).border = BORDA

# ------------------------------------------------------------------- aba Dados dos gráficos + Gráficos
ws_dg = wb.create_sheet("Dados dos gráficos")
ws_gr = wb.create_sheet("Gráficos")
ws_dg["A1"] = ("Dados que alimentam os gráficos. Os nomes estão na ordem do valor (maior primeiro) no momento da geração; os valores são fórmulas sobre a aba 'Estados'. "
               "Se os dados mudarem, a ordem pode ficar desatualizada.")
ws_dg["A1"].font, ws_dg["A1"].alignment = F_NORMAL, Alignment(wrap_text=True, vertical="top")
ws_dg.merge_cells("A1:H1")
ws_dg.row_dimensions[1].height = 44
E = "Estados!"
est_idx = tab_est.set_index("nome")
GRAFICOS = [
    ("saude_hab", "Gasto em Saúde por habitante, 2024 (R$)", "rs", False, "R$ por habitante"),
    ("saude_pct", "Saúde como % da despesa total, 2024", "pct", False, "% da despesa"),
    ("pess_pct", "Pessoal total (soma dos Poderes), % da RCL, 2024 (sem o DF)", "pct", True, "% da RCL ajustada"),
    ("dcl", "Dívida consolidada líquida, % da RCL, 2024", "pct", False, "% da RCL"),
]
valores_ordem = {
    "saude_hab": (tab_est.saude / tab_est["pop"]),
    "saude_pct": (tab_est.saude / tab_est.despesa02),
    "pess_pct": (tab_est[["pE", "pL", "pJ", "pM", "pD"]].sum(axis=1) / tab_est.rcl),
    "dcl": tab_est.dcl,
}
for g, (chave, titulo, fmt, sem_df, eixo) in enumerate(GRAFICOS):
    serie = pd.Series(valores_ordem[chave].values, index=tab_est.nome.values)
    if sem_df:
        serie = serie.drop("Distrito Federal")
    serie = serie.dropna().sort_values(ascending=False)
    c0 = 1 + g * 3
    ws_dg.cell(row=3, column=c0, value="Estado").font = F_CAB
    ws_dg.cell(row=3, column=c0 + 1, value=titulo.split(",")[0]).font = F_CAB
    for cc in (c0, c0 + 1):
        ws_dg.cell(row=3, column=cc).fill = FILL_CAB
        ws_dg.cell(row=3, column=cc).alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws_dg.row_dimensions[3].height = 48
    ws_dg.column_dimensions[LETRA(c0)].width = 26
    ws_dg.column_dimensions[LETRA(c0 + 1)].width = 16
    rj_idx = None
    for k, nome in enumerate(serie.index):
        r = 4 + k
        ws_dg.cell(row=r, column=c0, value=nome).font = F_NORMAL
        letra_col = let_est[chave if chave != "dcl" else "dcl"]
        f = f"=INDEX({E}${letra_col}$2:${letra_col}${ult_est},MATCH({LETRA(c0)}{r},{E}$A$2:$A${ult_est},0))"
        c = ws_dg.cell(row=r, column=c0 + 1, value=f)
        c.font, c.number_format = F_NORMAL, FMT[fmt]
        if nome == "Rio de Janeiro (estado)":
            rj_idx = k
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = titulo
    chart.legend = None
    chart.gapWidth = 40
    dados = Reference(ws_dg, min_col=c0 + 1, min_row=3, max_row=3 + len(serie))
    cats = Reference(ws_dg, min_col=c0, min_row=4, max_row=3 + len(serie))
    chart.add_data(dados, titles_from_data=True)
    chart.set_categories(cats)
    s = chart.series[0]
    s.graphicalProperties.solidFill = "C3C2B7"
    s.graphicalProperties.line.solidFill = "C3C2B7"
    if rj_idx is not None:
        pt = DataPoint(idx=rj_idx)
        pt.graphicalProperties.solidFill = "2A78D6"
        pt.graphicalProperties.line.solidFill = "2A78D6"
        s.dPt.append(pt)
    s.invertIfNegative = False                      # barras negativas mantêm a cor (sem isso ficam brancas)
    chart.title.overlay = False                     # o título não pode ficar por cima das barras
    chart.x_axis.scaling.orientation = "maxMin"
    chart.x_axis.tickLblPos = "low"                 # nomes dos estados à esquerda, mesmo com valores negativos
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.y_axis.crosses = "max"
    chart.y_axis.number_format = {"rs": '"R$" #,##0', "pct": "0%"}[fmt]
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill="E1E0D9", w=9525))   # grade clara e fina
    chart.width, chart.height = 17, 14              # a unidade já está no título; sem título de eixo
    ancora = ["B3", "L3", "B33", "L33"][g]
    ws_gr.add_chart(chart, ancora)
ws_gr["B1"] = "Gráficos dos 27 estados (barra azul = Rio de Janeiro; cinza = os demais). Dados em 'Dados dos gráficos'."
ws_gr["B1"].font = F_SUB

# ------------------------------------------------------------------------------------------ Legenda
ws_lg = wb.create_sheet("Legenda")
ws_lg.column_dimensions["A"].width = 26
ws_lg.column_dimensions["B"].width = 38
ws_lg.column_dimensions["C"].width = 70
ws_lg.column_dimensions["D"].width = 70


def titulo_lg(r, texto):
    ws_lg.cell(row=r, column=1, value=texto).font = F_SUB
    return r + 1


def cab_lg(r, itens):
    for i, h in enumerate(itens, start=1):
        c = ws_lg.cell(row=r, column=i, value=h)
        c.font, c.fill = F_CAB, FILL_CAB
    return r + 1


r = 1
ws_lg.cell(row=r, column=1, value="Legenda e glossário").font = F_TITULO
r += 2
r = titulo_lg(r, "1. Siglas e termos")
r = cab_lg(r, ["Sigla ou termo", "Significado"])
SIGLAS = [
    ("Siconfi", "Sistema de Informações Contábeis e Fiscais do Setor Público Brasileiro (Tesouro Nacional). Fonte de todos os números."),
    ("RREO", "Relatório Resumido da Execução Orçamentária (bimestral). Usamos o 6º bimestre, que acumula o ano."),
    ("RGF", "Relatório de Gestão Fiscal (quadrimestral; semestral nos municípios simplificados)."),
    ("LRF", "Lei de Responsabilidade Fiscal."),
    ("RCL", "Receita corrente líquida: base dos limites de pessoal e de endividamento."),
    ("RCL ajustada", "RCL menos transferências da União ligadas a emendas individuais, emendas de bancada e agentes comunitários de saúde. A dedução é feita pelo próprio ente no RGF."),
    ("Empenhada", "Despesa comprometida no orçamento. Não é o mesmo que paga."),
    ("Intra-orçamentárias", "Despesas entre órgãos do mesmo governo. Ficam de fora para não contar duas vezes."),
    ("Restos a pagar", "Despesas empenhadas e não pagas até o fim do ano."),
    ("Poder", "Executivo, Legislativo, Judiciário, Ministério Público ou Defensoria. A LRF fixa limite de pessoal para cada um."),
    ("Dívida consolidada líquida", "Dívida consolidada menos as disponibilidades de caixa e haveres, segundo o RGF."),
    ("Limite do Senado", "Limite de endividamento definido por resolução do Senado, como o ente o declara no RGF."),
    ("DF", "Distrito Federal. Acumula funções de estado e de município e só tem dado de 2 Poderes no RGF."),
    ("Posição (1 = maior)", "1 é o ente com o maior valor do indicador, dentro do grupo (27 estados, ou 92 municípios)."),
    ("n/c", "Não comparável: ente com Poderes faltando no RGF, fora da posição de pessoal total."),
    ("Mediana", "Valor do meio quando os entes são ordenados. Pouco sensível a casos extremos."),
]
for a, b in SIGLAS:
    ws_lg.cell(row=r, column=1, value=a).font = F_NEGRITO
    c = ws_lg.cell(row=r, column=2, value=b)
    c.font, c.alignment = F_NORMAL, Alignment(wrap_text=True, vertical="top")
    ws_lg.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    ws_lg.row_dimensions[r].height = 30
    r += 1
r += 1
r = titulo_lg(r, "2. Colunas das abas Estados e Municípios RJ")
r = cab_lg(r, ["Aba", "Coluna", "O que é", "De onde vem"])
for aba, cols in (("Estados", cols_est), ("Municípios RJ", cols_mun)):
    for c in cols:
        ws_lg.cell(row=r, column=1, value=aba).font = F_NORMAL
        ws_lg.cell(row=r, column=2, value=c[1]).font = F_NEGRITO
        x = ws_lg.cell(row=r, column=3, value=c[5])
        y = ws_lg.cell(row=r, column=4, value=c[6])
        for z in (x, y):
            z.font, z.alignment = F_NORMAL, Alignment(wrap_text=True, vertical="top")
        ws_lg.row_dimensions[r].height = 30
        r += 1
r += 1
r = titulo_lg(r, "3. Limites desta análise")
LIMITES = [
    "Um ano só (2024), e só o último período de cada relatório.",
    "Os valores são o que cada ente declarou. Não foram auditados.",
    "Limites de pessoal e de dívida são os declarados pelos entes, não conferidos com a lei.",
    "O Siconfi mostra gasto, não a qualidade do atendimento. Nenhum indicador aqui mede se a saúde de um lugar é boa ou ruim.",
    "O mínimo legal de saúde não está nesta planilha (fonte alternativa, o SIOPS, estava fora do ar em 21/10/2026).",
    "Comparações valem dentro de cada grupo: estados entre si, municípios entre si.",
    "Dados de 4 municípios estão marcados em amarelo por serem atípicos (RCL atípica ou totais de despesa inconsistentes). Estão mantidos, e a aba 'Com e sem' mostra o resultado sem eles.",
]
for texto_l in LIMITES:
    c = ws_lg.cell(row=r, column=1, value="• " + texto_l)
    c.font, c.alignment = F_NORMAL, Alignment(wrap_text=True, vertical="top")
    ws_lg.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
    ws_lg.row_dimensions[r].height = 28
    r += 1

# ------------------------------------------------------------------------------------------ Leia-me
ws_leia.column_dimensions["A"].width = 120
linhas = [
    ("Saúde fiscal do Rio de Janeiro: contas dos estados e dos municípios fluminenses, 2024", F_TITULO),
    ("", F_NORMAL),
    ("O que é esta planilha", F_SUB),
    ("Os dados públicos do Siconfi (Tesouro Nacional) usados no projeto, organizados para qualquer pessoa conferir: os valores de entrada, os indicadores calculados por fórmula, as posições do RJ entre os estados, gráficos e a legenda de cada coluna.", F_NORMAL),
    ("Pergunta do projeto: como o Rio de Janeiro e os seus municípios estão gerindo o dinheiro da saúde, e onde as contas públicas estão apertadas?", F_NORMAL),
    ("", F_NORMAL),
    ("Como ler", F_SUB),
    ("• Texto AZUL = dado que veio da API do Siconfi (entrada). Texto PRETO = fórmula: clique na célula para ver o cálculo.", F_NORMAL),
    ("• Linha azul clara = Rio de Janeiro (o estado, na aba Estados; a capital, na aba Municípios RJ). Linha amarela = dado atípico, mantido e marcado.", F_NORMAL),
    ("• Os indicadores são fórmulas sobre as colunas de entrada, então qualquer número pode ser refeito à mão.", F_NORMAL),
    ("", F_NORMAL),
    ("Abas", F_SUB),
    ("• Estados: os 27 estados (com o DF), com entradas, indicadores e posições.", F_NORMAL),
    ("• Municípios RJ: os 92 municípios do RJ, com os mesmos campos e a marca dos 12 que entregam relatório simplificado.", F_NORMAL),
    ("• Com e sem: medianas dos municípios com e sem os casos especiais.", F_NORMAL),
    ("• Gráficos e Dados dos gráficos: barras dos estados, com o RJ em azul.", F_NORMAL),
    ("• Legenda: siglas, descrição e origem de cada coluna, e limites da análise.", F_NORMAL),
    ("", F_NORMAL),
    ("De onde vêm os dados", F_SUB),
    ("• Fonte: API de dados abertos do Tesouro Nacional (Siconfi), https://apidatalake.tesouro.gov.br/ords/siconfi/tt/ . Documentação: https://apidatalake.tesouro.gov.br/docs/siconfi/", F_NORMAL),
    ("• Exercício: 2024. Coleta feita em 21/10/2026.", F_NORMAL),
    ("• Relatórios usados: RREO (Anexos 01, 02 e 07, 6º bimestre) e RGF (Anexo 01, por Poder, e Anexo 02, no último período entregue por cada ente).", F_NORMAL),
    ("• Os números são o que cada ente declarou ao Tesouro. Não foram auditados.", F_NORMAL),
    ("", F_NORMAL),
    ("Como foi feito (reprodutível)", F_SUB),
    ("• Coleta: scripts/04_coleta_valores.py. Conferência: scripts/05_conferencia_coleta.py. EDA: notebooks/01_eda_siconfi_2024.ipynb. Esta planilha: scripts/09_gerar_planilha.py.", F_NORMAL),
    ("• Registro de fontes e de pendências: docs/02-fontes-de-dados.md e docs/04-pendencias-measure.md.", F_NORMAL),
    ("", F_NORMAL),
    ("Cuidados ao interpretar", F_SUB),
    ("• O Siconfi mostra quanto se gasta, e não a qualidade do atendimento. Esta planilha não diz qual lugar tem melhor ou pior saúde.", F_NORMAL),
    ("• Compare estados com estados e municípios com municípios. Os dois grupos têm papéis e receitas diferentes na saúde.", F_NORMAL),
    ("• Saúde por habitante acompanha de perto a receita por habitante (correlação de 0,80 nos estados e 0,82 nos municípios). Para a prioridade dada à saúde, prefira 'Saúde, % da despesa'.", F_NORMAL),
    ("• O indicador 'Saúde, % da despesa' NÃO é o mínimo legal de saúde, que usa outra base (receita de impostos e transferências).", F_NORMAL),
    ("• Os limites de pessoal e de dívida são os que os entes declararam, ainda não conferidos com a lei.", F_NORMAL),
    ("• Dívida e restos a pagar são mostrados de forma descritiva. Nenhuma hipótese foi testada nesta planilha.", F_NORMAL),
]
for i, (texto_l, fonte) in enumerate(linhas, start=1):
    c = ws_leia.cell(row=i, column=1, value=texto_l)
    c.font, c.alignment = fonte, Alignment(wrap_text=True, vertical="top")
    if fonte is F_NORMAL and len(texto_l) > 120:
        ws_leia.row_dimensions[i].height = 15 * (1 + len(texto_l) // 118)
wb.properties.title = "Saúde fiscal do RJ, 2024"
wb.properties.creator = "Projeto saude-fiscal-rj"

SAIDA.parent.mkdir(exist_ok=True)
wb.save(SAIDA)
print("Planilha gravada:", SAIDA)
print("Estados:", len(tab_est), "| Municípios:", len(tab_mun), "| colunas estados:", len(cols_est), "| colunas municípios:", len(cols_mun))

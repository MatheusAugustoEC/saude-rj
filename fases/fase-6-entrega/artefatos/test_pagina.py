"""Testes da página de entrega: os números da página são os do pipeline, e os arquivos que ela oferece existem e abrem.

Uso: python -m pytest tests -q   (depois de rodar scripts/35 e scripts/36)
"""
import hashlib
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parent.parent
PAG = RAIZ / "pagina"
DADOS = json.loads((RAIZ / "dados_tratados" / "pagina_dados.json").read_text(encoding="utf-8"))
HTML = (PAG / "index.html").read_text(encoding="utf-8")


def test_dados_batem_com_o_pipeline():
    ind = pd.read_csv(RAIZ / "dados_tratados" / "indicadores_2024.csv")
    mun = ind[ind.esfera == "M"]
    assert len(DADOS["cidades"]) == len(mun) == 92 and len(DADOS["ufs"]) == 27
    g = {c["c"]: c["gh"] for c in DADOS["cidades"]}
    for r in mun.itertuples():
        assert abs(g[int(r.cod_ibge)] - r.saude_por_hab) < 0.06
    rj = [u for u in DADOS["ufs"] if u["s"] == "RJ"][0]
    assert rj["gh"] == min(u["gh"] for u in DADOS["ufs"]) and DADOS["estado_rj"]["pos_gh"] == 27


def test_numeros_chave():
    g = DADOS["globais"]
    assert g["pares_reducao"] == 46 and (g["k_sustentam"], g["k_total"]) == (8, 9)
    assert (g["folga_menor1"], g["folga_menor3"], g["acima_prudencial"], g["siops_fora"]) == (7, 18, 22, 11)
    assert round(g["razao_extremos"]) == 7 and 0.79 < g["elast_total"]["2024"] < 0.81


def test_sem_valor_invalido_nos_dados():
    texto = json.dumps(DADOS, ensure_ascii=False)
    assert "NaN" not in texto and "Infinity" not in texto
    assert all(c["gh"] and c["rh"] and c["ap"] is not None for c in DADOS["cidades"])


def test_toda_curiosidade_tem_fonte_e_janela():
    for c in DADOS["curiosidades"]:
        assert c["j"] and (RAIZ / c["f"]).exists(), c["ti"]


def test_malhas_batem_com_o_hash():
    for nome, h in DADOS["meta"]["malhas_sha256"].items():
        assert hashlib.sha256((RAIZ / "dados_brutos" / "malha_ibge_api" / nome).read_bytes()).hexdigest() == h


def test_arquivos_da_biblioteca_existem():
    bib = json.loads(re.search(r"const BIB=(\[.*?\]);", HTML, re.S).group(1))
    assert len(bib) >= 20
    for b in bib:
        if not b.get("ext"):
            assert (PAG / b["arq"]).exists(), b["arq"]


def test_zip_e_planilha_abrem():
    z = zipfile.ZipFile(PAG / "arquivos" / "saude-fiscal-rj-biblioteca.zip")
    assert z.testzip() is None and len(z.namelist()) == len(set(z.namelist())) >= 20
    wb = openpyxl.load_workbook(PAG / "arquivos" / "saude-fiscal-rj-cidades.xlsx")
    assert wb.sheetnames == ["Sobre", "Cidades", "Regiões", "Estados"]
    assert wb["Cidades"].max_row == 93 and wb["Estados"].max_row == 28


def test_pagina_sem_cdn_nem_marcador_pendente():
    assert not re.search(r"<script[^>]+src=\"https?://", HTML) and "/*CSS*/" not in HTML and "lorem" not in HTML.lower()
    assert (PAG / "pptxgen.bundle.js").exists()
